import docx
import os

doc_path = 'Documentation/NexusIVR_Documentation.docx'
doc = docx.Document(doc_path)

print("Step 1: Replacing diagram images in relationship parts...")

image_replacements = {
    'media/image8.png': 'scratch/generated_diagrams/fig_4_2_system_architecture.png',   # Figure 4.2
    'media/image9.png': 'scratch/generated_diagrams/fig_6_1_7pass_pipeline.png',          # Figure 6.1
    'media/image16.png': 'scratch/generated_diagrams/fig_11_1_erd.png',                   # Figure 11.1
    'media/image17.png': 'scratch/generated_diagrams/fig_12_1_multitenant.png',           # Figure 12.1
    'media/image24.png': 'scratch/generated_diagrams/fig_startup_containers.png'          # Container Startup
}

replaced_count = 0
for rel in doc.part.rels.values():
    if 'image' in rel.target_ref:
        target_ref = rel.target_ref
        if target_ref in image_replacements:
            new_img_path = image_replacements[target_ref]
            if os.path.exists(new_img_path):
                with open(new_img_path, 'rb') as f:
                    new_blob = f.read()
                rel.target_part._blob = new_blob
                replaced_count += 1
                print(f"  Updated image: {target_ref} -> {new_img_path} ({len(new_blob)} bytes)")

print(f"Total images updated: {replaced_count}")

print("\nStep 2: Updating text in Section 11.3 Vector Search & RAG Pipeline...")

# Find Section 11.3 heading in body (P280)
target_idx = None
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == "11.3 Vector Search & RAG Pipeline" and i > 100:
        target_idx = i
        break

if target_idx is not None:
    print(f"Found Section 11.3 at paragraph index {target_idx}")

    # Section 11.3 content paragraphs to insert
    new_section_paragraphs = [
        ("Heading 3", "11.3.1 Architectural Overview & Core Intent"),
        ("Normal", "The Retrieval-Augmented Generation (RAG) pipeline equips the NexusIVR AI Assistant with deep, domain-specific knowledge about a tenant company's business rules, product catalogs, service pricing, and operational policies. Rather than relying solely on the static parametric knowledge of commercial LLMs (which leads to generic or hallucinatory answers), the RAG pipeline grounds conversational turns and IVR generation requests in authoritative tenant documentation uploaded by Tenant Administrators."),
        ("Normal", "When a user submits a prompt or chat question, KnowledgeService queries the tenant's vector database to retrieve the top-K most semantically relevant document chunks. These chunks are dynamically injected into the LLM system prompt as grounding context, enforcing strict factual alignment with company documentation and enabling inline source citations."),
        
        ("Heading 3", "11.3.2 Document Ingestion, Text Chunking & Embedding Pipeline"),
        ("Normal", "1. Document Upload & Parsing: Tenant Admins upload PDF, DOCX, TXT, VXML, CSV, or Markdown documents via the UI (/api/v1/ai/rag/upload). KnowledgeService parses raw text and extracts structural metadata (filename, page numbers, section headers, uploaded_at)."),
        ("Normal", "2. Text Chunking Strategy: Documents are split using a sliding window chunking algorithm with a target size of 512 tokens and a 64-token overlap. The overlap ensures continuity and prevents semantic truncation across sentence boundaries."),
        ("Normal", "3. Vector Embedding Generation: Text chunks are transformed into 1536-dimensional dense vector embeddings using OpenAI's text-embedding-3-small or custom vector embedding models. Embeddings are L2-normalized prior to persistence in PostgreSQL."),

        ("Heading 3", "11.3.3 Database Vector Schema & HNSW Indexing Strategy"),
        ("Normal", "The RAG storage architecture utilizes PostgreSQL 15+ with the pgvector extension enabled. Document metadata is stored in knowledge_docs, and vector chunks are stored in knowledge_embeddings."),
        ("Normal", "CREATE TABLE knowledge_docs (\n"
                   "    doc_id UUID PRIMARY KEY DEFAULT gen_random_uuid(),\n"
                   "    tenant_id UUID NOT NULL REFERENCES tenants(id) ON DELETE CASCADE,\n"
                   "    filename VARCHAR(255) NOT NULL,\n"
                   "    file_type VARCHAR(50) NOT NULL,\n"
                   "    status VARCHAR(50) DEFAULT 'PROCESSING',\n"
                   "    created_at TIMESTAMP WITH TIME ZONE DEFAULT CURRENT_TIMESTAMP\n"
                   ");"),
        ("Normal", "CREATE TABLE knowledge_embeddings (\n"
                   "    embedding_id UUID PRIMARY KEY DEFAULT gen_random_uuid(),\n"
                   "    doc_id UUID NOT NULL REFERENCES knowledge_docs(doc_id) ON DELETE CASCADE,\n"
                   "    tenant_id UUID NOT NULL REFERENCES tenants(id) ON DELETE CASCADE,\n"
                   "    chunk_index INT NOT NULL,\n"
                   "    chunk_content TEXT NOT NULL,\n"
                   "    embedding vector(1536) NOT NULL,\n"
                   "    created_at TIMESTAMP WITH TIME ZONE DEFAULT CURRENT_TIMESTAMP\n"
                   ");\n\n"
                   "CREATE INDEX idx_embeddings_vector_hnsw \n"
                   "ON knowledge_embeddings \n"
                   "USING hnsw (embedding vector_cosine_ops) \n"
                   "WITH (m = 16, ef_construction = 64);"),
        ("Normal", "The Hierarchical Navigable Small World (HNSW) graph index enables sub-10ms approximate nearest neighbor (ANN) vector similarity search across hundreds of thousands of document chunks. Vector distance queries utilize pgvector's cosine distance operator (<=>) filtered by tenant_id:"),
        ("Normal", "SELECT chunk_content, doc_id, (embedding <=> ?::vector) AS distance\n"
                   "FROM knowledge_embeddings\n"
                   "WHERE tenant_id = ?\n"
                   "ORDER BY embedding <=> ?::vector ASC\n"
                   "LIMIT 5;"),

        ("Heading 3", "11.3.4 Python RAG Microservice & Java RagClient Integration"),
        ("Normal", "To ensure modular separation of concerns and high throughput, vector retrieval is encapsulated within a Python FastAPI microservice (rag-service listening on port 8000), which interfaces with IVR-AI-engine via RagClient.java."),
        ("Normal", "• API Request: POST http://localhost:8000/query\n"
                   "  Payload: {\"query\": \"What are your business hours?\", \"top_k\": 5, \"min_score\": 0.35, \"tenant_id\": \"...\"}\n"
                   "• Response: Returns matched citations list, max similarity score, and fallback_required boolean flag.\n"
                   "• Fallback Protection: If rag-service is unreachable or returns zero matching chunks above threshold (0.35), RagClient sets fallbackRequired=true. ChatService gracefully proceeds with ungrounded LLM generation without throwing user-facing errors."),

        ("Heading 3", "11.3.5 Grounded Prompt Engineering & Citation Generation"),
        ("Normal", "When matching knowledge chunks are retrieved by RagClient, ChatService constructs a grounded prompt block:"),
        ("Normal", "### Relevant Context:\n"
                   "- [Source: company_policy.pdf, Page 3] Our business hours are Monday through Friday, 9 AM to 6 PM EST.\n"
                   "- [Source: faq_catalog.docx, Section 2] Weekend emergency support is available for Premium tiers.\n\n"
                   "### Instructions:\n"
                   "Use ONLY the Relevant Context provided above to answer the user's question. "
                   "Cite sources inline using [Source: filename, Section/Page]. "
                   "If the answer cannot be determined from the context, state explicitly: "
                   "\"I cannot find information about this in the project documentation.\""),
        ("Normal", "Inline source citations are parsed and returned in the REST API response as SourceCitation DTO objects, rendered as interactive badge links in the web UI."),

        ("Heading 3", "11.3.6 Multi-Tenant Isolation & Security for Vector Data"),
        ("Normal", "Vector similarity searches strictly enforce multi-tenant boundaries at all three security layers:\n"
                   "1. Schema Layer: Mandatory tenant_id column on knowledge_docs and knowledge_embeddings tables.\n"
                   "2. Application Layer: BaseAiServlet extracts tenant_id from the verified JWT token, and RagClient injects mandatory tenant_id filtering into every vector search query.\n"
                   "3. Database RLS Layer: PostgreSQL Row-Level Security policies restrict vector distance searches (<=>) to the authenticated tenant_id, preventing cross-tenant document leakage.")
    ]

    # Replace existing paragraphs P281, P282, P283 with the expanded subsections
    p_281 = doc.paragraphs[target_idx + 1]
    p_282 = doc.paragraphs[target_idx + 2]
    p_283 = doc.paragraphs[target_idx + 3]

    # Update p_281 to first new subsection heading
    p_281.style = doc.styles['Heading 3']
    p_281.text = new_section_paragraphs[0][1]

    # Update p_282 and p_283 to text
    p_282.style = doc.styles['Normal']
    p_282.text = new_section_paragraphs[1][1]

    p_283.style = doc.styles['Normal']
    p_283.text = new_section_paragraphs[2][1]

    # Insert remaining paragraphs after p_283
    ref_p = p_283
    for style_name, text_content in new_section_paragraphs[3:]:
        new_p = doc.add_paragraph(text_content, style=style_name)
        ref_p._element.addnext(new_p._element)
        ref_p = new_p

    print("Successfully expanded Section 11.3 in docx!")

# Save updated document
doc.save(doc_path)
print(f"\nSaved updated document to: {doc_path}")
